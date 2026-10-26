"""文件解析：把上传文件提取成纯文本

MVP 支持：TXT / PDF / DOCX
扩展预留：PPT（python-pptx）、图片 OCR
"""
from pathlib import Path

# 解析出的文本上限（字符），避免超大文件占满内存
MAX_TEXT_LENGTH = 200_000


def parse_file(path: Path, file_type: str) -> str:
    """按文件类型解析，返回纯文本；失败抛出 ValueError"""
    return parse_document(path, file_type)[0]


def parse_document(path: Path, file_type: str) -> tuple[str, list[dict]]:
    """保留真实页码/段落/行号；偏移始终相对最终保存的解析文本。"""
    if file_type == "txt":
        text = _parse_txt(path)
        left = len(text) - len(text.lstrip())
        parts, offset = [], left
        first_line = text[:left].count('\n') + 1
        for line_number, line in enumerate(text[left:left+MAX_TEXT_LENGTH].splitlines(keepends=True), first_line):
            end = offset + len(line.rstrip('\r\n'))
            parts.append(dict(start=offset, end=end, kind='line', number=line_number, label=f'第 {line_number} 行'))
            offset += len(line)
    elif file_type == "pdf":
        from pypdf import PdfReader
        pages = [(page.extract_text() or '') for page in PdfReader(str(path)).pages]
        text, parts = _join_parts([(body, 'page', index, f'第 {index} 页') for index, body in enumerate(pages, 1)])
    elif file_type == "docx":
        from docx import Document
        document = Document(str(path))
        chunks = [(p.text, 'paragraph', index, f'第 {index} 段') for index, p in enumerate(document.paragraphs, 1)]
        for table_number, table in enumerate(document.tables, 1):
            for row_number, row in enumerate(table.rows, 1):
                chunks.append((' | '.join(cell.text for cell in row.cells), 'table_row', row_number, f'表 {table_number} · 第 {row_number} 行'))
        text, parts = _join_parts(chunks)
    else:
        raise ValueError(f"暂不支持的文件类型：{file_type}")
    left = len(text) - len(text.lstrip())
    retained = text.strip()[:MAX_TEXT_LENGTH]
    locations = []
    for part in parts:
        start, end = max(0, part['start'] - left), min(len(retained), part['end'] - left)
        if start < end and retained[start:end].strip():
            locations.append({**part, 'start': start, 'end': end})
    return retained, locations


def _join_parts(chunks):
    texts, locations, offset = [], [], 0
    for text, kind, number, label in chunks:
        texts.append(text)
        locations.append(dict(start=offset, end=offset + len(text), kind=kind, number=number, label=label))
        offset += len(text) + 1
    return '\n'.join(texts), locations


def _parse_txt(path: Path) -> str:
    # 常见编码依次尝试：UTF-8 / GBK / UTF-16
    for encoding in ("utf-8", "gbk", "utf-16"):
        try:
            return path.read_text(encoding=encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
    # 全部失败时忽略错误字符兜底
    return path.read_text(encoding="utf-8", errors="ignore")
