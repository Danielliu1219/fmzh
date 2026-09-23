# -*- coding: utf-8 -*-
"""PDF / DOCX 解析验证：生成真实文件 -> 上传 -> 检查解析结果

用法：后端启动后执行  .venv/Scripts/python.exe verify_formats.py
"""
import io
import sys

import httpx

BASE = "http://127.0.0.1:8000"


def make_pdf(text: str) -> bytes:
    """手工构造一个最简单的 PDF（Helvetica 字体，ASCII 文本）"""
    content = f"BT /F1 18 Tf 72 700 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n%s\nendobj\n" % (i, obj))
    xref_pos = out.tell()
    out.write(b"xref\n0 %d\n" % (len(objects) + 1))
    out.write(b"0000000000 65535 f \n")
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
        % (len(objects) + 1, xref_pos)
    )
    return out.getvalue()


def make_docx(text: str) -> bytes:
    """用 python-docx 生成真实 Word 文档"""
    from docx import Document

    doc = Document()
    doc.add_heading("高等数学复习重点", level=1)
    doc.add_paragraph(text)
    doc.add_paragraph("必考：洛必达法则、泰勒公式。")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


client = httpx.Client(timeout=60)

# 登录（复用 smoke_test 创建的账号；没有就注册）
r = client.post(f"{BASE}/api/auth/login", json={"username": "testuser", "password": "test123456"})
if r.status_code != 200:
    r = client.post(
        f"{BASE}/api/auth/register",
        json={"username": "testuser", "password": "test123456", "nickname": "测试同学"},
    )
token = r.json()["access_token"]
headers = {"Authorization": f"Bearer {token}"}

r = client.post(f"{BASE}/api/exams", headers=headers, json={"subject": "线性代数", "exam_date": "2026-09-15"})
exam_id = r.json()["id"]

# 上传 PDF
r = client.post(
    f"{BASE}/api/materials",
    headers=headers,
    files={"file": ("线代重点.pdf", io.BytesIO(make_pdf("Matrix eigenvector determinant SVD")), "application/pdf")},
    data={"title": "线代重点PDF", "exam_id": str(exam_id), "source_type": "老师重点"},
)
print("PDF 上传:", r.status_code, "parse_status =", r.json().get("parse_status"))
detail = client.get(f"{BASE}/api/materials/{r.json()['id']}", headers=headers).json()
assert r.json()["parse_status"] == "done" and "Matrix" in detail["text_preview"], "PDF 解析失败"
print("  [PASS] PDF 解析出文本")

# 上传 DOCX
r = client.post(
    f"{BASE}/api/materials",
    headers=headers,
    files={"file": ("线代笔记.docx", io.BytesIO(make_docx("矩阵的秩与线性方程组解的关系，重点掌握行列式展开。")), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    data={"title": "线代笔记DOCX", "exam_id": str(exam_id), "source_type": "课堂笔记"},
)
print("DOCX 上传:", r.status_code, "parse_status =", r.json().get("parse_status"))
detail = client.get(f"{BASE}/api/materials/{r.json()['id']}", headers=headers).json()
assert r.json()["parse_status"] == "done" and "线性方程组" in detail["text_preview"], "DOCX 解析失败"
print("  [PASS] DOCX 解析出文本（含中文）")

# 该科目资料齐了，顺带验证分析可用
r = client.post(
    f"{BASE}/api/ai/profile",
    headers=headers,
    json={"exam_id": exam_id, "raw_text": "线代基础还行，想冲高分，还有8天考试，每天能学4小时"},
)
assert r.json()["is_complete"], "画像应完整：" + str(r.json())
r = client.post(f"{BASE}/api/ai/analyze", headers=headers, json={"exam_id": exam_id})
assert r.status_code == 200, r.text
print("  [PASS] 该科目完整走通画像 + 分析")

print("\nPDF/DOCX 验证全部通过")
sys.exit(0)
