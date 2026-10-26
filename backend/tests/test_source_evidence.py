import json
import unittest
from datetime import date, timedelta
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from docx import Document
from fastapi import Depends, FastAPI, Header
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.config import settings
from app.database import Base, get_db
from app.models.models import Analysis, Exam, Material, Task, User, UserProfile
from app.routers import ai, materials
from app.schemas.schemas import AnalysisResult, KnowledgePoint, SourceReference
from app.services.file_parser import MAX_TEXT_LENGTH, parse_document, parse_file
from app.services.schema_upgrade import upgrade_schema
from app.services.source_evidence import attach_references, content_hash, excerpt_for, references_for
from app.utils.security import get_current_user


BODY = '第一章\n积分是求和的极限，是微积分的重要概念。\n第二章\n导数描述函数的变化率。'


class EvidenceTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory(); self.root = Path(self.temp.name)
        self.engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
        Base.metadata.create_all(self.engine)
        self.db = Session(self.engine)
        self.owner = User(username='owner', hashed_password='unused')
        self.peer = User(username='peer', hashed_password='unused')
        self.db.add_all([self.owner, self.peer]); self.db.flush()
        self.exam = Exam(user_id=self.owner.id, subject='数学', exam_date=date.today()+timedelta(days=6))
        self.other = Exam(user_id=self.owner.id, subject='英语', exam_date=date.today()+timedelta(days=8))
        self.peer_exam = Exam(user_id=self.peer.id, subject='数学', exam_date=date.today()+timedelta(days=6))
        self.db.add_all([self.exam, self.other, self.peer_exam]); self.db.flush()
        self.profile = UserProfile(user_id=self.owner.id, exam_id=self.exam.id, target='pass', target_band='D', learning_status='weak', available_days=4, daily_hours=2, is_complete=True)
        self.db.add(self.profile); self.db.commit()
        app = FastAPI()
        def database():
            with Session(self.engine) as db: yield db
        def current(x_user_id: int = Header(default=self.owner.id), db=Depends(get_db)):
            return db.get(User, x_user_id)
        app.dependency_overrides[get_db] = database
        app.dependency_overrides[get_current_user] = current
        app.include_router(ai.router, prefix='/api/ai')
        app.include_router(materials.router, prefix='/api/materials')
        self.client = TestClient(app, raise_server_exceptions=False)

    def tearDown(self):
        self.client.close(); self.db.close(); self.engine.dispose(); self.temp.cleanup()

    def material(self, body=BODY, owner=None, exam=None, locations=None):
        row = Material(user_id=(owner or self.owner).id, exam_id=(exam or self.exam).id, title='积分讲义',
                       source_type='老师重点', file_path=str(self.root/'test.txt'), file_type='txt', text_content=body,
                       text_locations=json.dumps(locations or []), parse_status='done')
        self.db.add(row); self.db.commit()
        return row

    def analysis(self, material=None, references=None, name='积分'):
        result = AnalysisResult(course='数学', knowledge_points=[KnowledgePoint(name=name, weight=0.8, source='老师重点', citations=references or [])],
                                quiz=[dict(kind='fill', question='积分题', answer='42', knowledge_point='积分')])
        if material: attach_references(result, [material])
        row = Analysis(user_id=self.owner.id, exam_id=self.exam.id, profile_id=self.profile.id, result_json=result.model_dump_json())
        self.db.add(row); self.db.commit()
        return row

    def lookup(self, analysis, status=200, kind='knowledge', index=0, headers=None):
        response = self.client.post(f'/api/ai/analyses/{self.exam.id}/evidence', json=dict(analysis_id=analysis.id, kind=kind, item_index=index), headers=headers)
        self.assertEqual(response.status_code, status, response.text)
        return response.json()

    def test_quoted_citation_ignores_model_page_and_offsets_uses_exact_original(self):
        material = self.material(locations=[dict(start=4, end=24, kind='page', number=2, label='第 2 页')])
        proposed = SourceReference(material_id=material.id, quote='积分是求和的极限，是微积分的重要概念。', start=999, end=1000, content_hash='forged')
        analysis = self.analysis(material, [proposed])
        item = self.lookup(analysis)['items'][0]
        self.assertEqual(item['status'], 'available'); self.assertEqual(item['page'], 2)
        self.assertEqual(BODY[item['start']:item['end']], proposed.quote)
        self.assertEqual(item['context'][item['highlight_start']:item['highlight_end']], proposed.quote)

    def test_quote_with_pdf_line_wrap_maps_back_to_original_characters(self):
        material = self.material(body='Intro\nIntegral\n is a limit of sums.\nEnd')
        point = KnowledgePoint(name='积分', weight=.8, source='老师重点', citations=[SourceReference(material_id=material.id, quote='Integral is a limit of sums.')])
        reference = references_for(point, [material])[0]
        self.assertEqual(reference.quote, 'Integral\n is a limit of sums.')
        self.assertEqual(reference.match, 'quote')

    def test_forged_quote_and_unrelated_keywords_produce_no_invented_reference(self):
        material = self.material()
        analysis = self.analysis(material, [SourceReference(material_id=material.id, quote='完全不存在的引文内容')], name='不存在的知识点')
        response = self.lookup(analysis)
        self.assertEqual(response['items'], [])
        self.assertIn('未找到', response['note'])

    def test_legacy_analysis_gets_explicit_keyword_match_without_fake_page_or_writes(self):
        self.material()
        analysis = self.analysis()
        before = analysis.result_json
        item = self.lookup(analysis)['items'][0]
        self.assertEqual(item['match'], 'keyword'); self.assertIsNone(item['page'])
        self.assertIn('旧资料未保存页码', item['location'])
        self.assertIn('不代表已核实', item['note'])
        self.db.expire_all(); self.assertEqual(self.db.get(Analysis, analysis.id).result_json, before)
        self.assertEqual(self.db.query(Task).count(), 0)

    def test_cross_user_and_cross_subject_materials_never_appear(self):
        own = self.material()
        peer = self.material(body='secret-peer-body', owner=self.peer, exam=self.peer_exam)
        other = self.material(body='secret-other-subject', exam=self.other)
        analysis = self.analysis(own, [SourceReference(material_id=peer.id, quote='secret-peer-body'), SourceReference(material_id=other.id, quote='secret-other-subject')])
        response = self.lookup(analysis)
        self.assertNotIn('secret-', json.dumps(response))
        self.assertTrue(all(item['material_id'] == own.id for item in response['items']))
        self.lookup(analysis, 404, headers={'X-User-ID': str(self.peer.id)})

    def test_changed_content_invalidates_snapshot_even_if_keyword_remains(self):
        material = self.material(); analysis = self.analysis(material)
        material.text_content = BODY + '\n资料已更新'; self.db.commit()
        item = self.lookup(analysis)['items'][0]
        self.assertEqual(item['status'], 'changed'); self.assertEqual(item['context'], '')

    def test_deleted_and_moved_materials_are_unavailable(self):
        material = self.material(); analysis = self.analysis(material)
        material.exam_id = self.other.id; self.db.commit()
        item = self.lookup(analysis)['items'][0]
        self.assertEqual(item['status'], 'unavailable'); self.assertEqual(item['context'], '')
        self.db.delete(material); self.db.commit()
        self.assertEqual(self.lookup(analysis)['items'][0]['status'], 'unavailable')

    def test_old_analysis_id_is_used_even_after_new_analysis_exists(self):
        material = self.material(); older = self.analysis(material)
        self.analysis(material, name='导数')
        self.assertEqual(self.lookup(older)['item_title'], '积分')

    def test_foreign_analysis_and_invalid_item_are_rejected(self):
        analysis = self.analysis()
        self.lookup(analysis, 404, index=99)
        self.lookup(analysis, 422, kind='unsupported')
        response = self.client.post(f'/api/ai/analyses/{self.other.id}/evidence', json=dict(analysis_id=analysis.id, kind='knowledge', item_index=0))
        self.assertEqual(response.status_code, 404)

    def test_material_excerpt_route_handles_deep_offsets_and_stale_or_foreign_links(self):
        material = self.material(body='a'*1500+'积分原文'+'b'*300)
        params = dict(start=1500, end=1504, content_hash=content_hash(material.text_content), material_created_at=material.created_at.isoformat())
        response = self.client.get(f'/api/materials/{material.id}/excerpt', params=params)
        self.assertEqual(response.status_code, 200, response.text)
        item = response.json()
        self.assertEqual(item['context'][item['highlight_start']:item['highlight_end']], '积分原文')
        self.assertEqual(self.client.get(f'/api/materials/{material.id}/excerpt', params=params, headers={'X-User-ID': str(self.peer.id)}).status_code, 404)
        material.text_content += 'changed'; self.db.commit()
        self.assertEqual(self.client.get(f'/api/materials/{material.id}/excerpt', params=params).status_code, 409)
        self.assertEqual(self.client.get(f'/api/materials/{material.id}/excerpt', params={**params, 'end': 99999}).status_code, 416)

    def test_reused_material_id_does_not_inherit_old_reference(self):
        material = self.material(); analysis = self.analysis(material)
        old_id, created = material.id, material.created_at
        self.db.delete(material); self.db.commit()
        replacement = Material(id=old_id, user_id=self.owner.id, exam_id=self.exam.id, title='替换后的资料',
                               file_path='unused', file_type='txt', text_content=BODY, created_at=created+timedelta(seconds=1))
        self.db.add(replacement); self.db.commit()
        self.assertEqual(self.lookup(analysis)['items'][0]['status'], 'unavailable')

    def test_new_analysis_saves_verified_references_and_sends_material_ids(self):
        material = self.material()
        model_result = dict(course='数学', knowledge_points=[dict(name='积分', weight=.8, source='老师重点', citations=[dict(material_id=material.id, quote='积分是求和的极限，是微积分的重要概念。')])],
                            quiz=[dict(kind='fill', question='积分题', answer='42', knowledge_point='积分')])
        with patch.object(settings, 'LLM_MOCK', False), patch.object(ai, 'chat_json', return_value=model_result) as model:
            response = self.client.post('/api/ai/analyze', json={'exam_id': self.exam.id})
        self.assertEqual(response.status_code, 200, response.text)
        reference = response.json()['result']['knowledge_points'][0]['citations'][0]
        self.assertEqual(reference['content_hash'], content_hash(BODY))
        self.assertIn(f'资料编号 {material.id}', model.call_args.args[1])
        self.assertEqual(self.lookup(self.db.get(Analysis, response.json()['id']))['items'][0]['status'], 'available')

    def test_txt_line_numbers_and_trim_offsets_survive_truncation(self):
        path = self.root/'lines.txt'; path.write_text('\n\n  开头\n积分原文\n'+'x'*MAX_TEXT_LENGTH, encoding='utf-8')
        body, locations = parse_document(path, 'txt')
        self.assertEqual(len(body), MAX_TEXT_LENGTH)
        line = next(item for item in locations if item['number'] == 4)
        self.assertEqual(body[line['start']:line['end']], '积分原文')
        self.assertEqual(parse_file(path, 'txt'), body)

    def test_docx_paragraph_and_table_locations_are_real(self):
        path = self.root/'paragraphs.docx'; document = Document()
        document.add_paragraph('第一段'); document.add_paragraph('积分原文')
        table = document.add_table(rows=1, cols=2); table.cell(0,0).text='表格'; table.cell(0,1).text='依据'
        document.save(path)
        body, locations = parse_document(path, 'docx')
        paragraph = next(item for item in locations if item['kind']=='paragraph' and item['number']==2)
        self.assertEqual(body[paragraph['start']:paragraph['end']], '积分原文')
        self.assertEqual(locations[-1]['label'], '表 1 · 第 1 行')

    def test_pdf_page_numbers_and_original_offsets(self):
        path = self.root/'pages.pdf'; writer = PdfWriter()
        font = DictionaryObject({NameObject('/Type'):NameObject('/Font'), NameObject('/Subtype'):NameObject('/Type1'), NameObject('/BaseFont'):NameObject('/Helvetica')})
        for value in ['First page', 'Integral evidence on page two']:
            page = writer.add_blank_page(width=300, height=200)
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
            stream = DecodedStreamObject(); stream.set_data(f'BT /F1 12 Tf 20 100 Td ({value}) Tj ET'.encode())
            page[NameObject('/Contents')] = writer._add_object(stream)
        writer.write(path)
        body, locations = parse_document(path, 'pdf')
        material = self.material(body, locations=locations); material.file_type='pdf'; self.db.commit()
        reference = references_for(KnowledgePoint(name='Integral', weight=.8), [material])[0]
        self.assertEqual(excerpt_for(material, reference)['page'], 2)

    def test_upload_persists_locations_and_failed_parse_keeps_empty_locations(self):
        with patch.object(materials, 'UPLOAD_DIR', self.root):
            response = self.client.post('/api/materials', data={'title':'new','exam_id':self.exam.id}, files={'file':('note.txt', BODY.encode(), 'text/plain')})
            self.assertEqual(response.status_code, 200, response.text)
            row=self.db.get(Material, response.json()['id']); self.assertTrue(json.loads(row.text_locations))
            response=self.client.post('/api/materials', data={'title':'invalid'}, files={'file':('bad.pdf', b'%PDF-1.4\ninvalid', 'application/pdf')})
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()['parse_status'], 'failed')
            self.assertEqual(self.db.get(Material,response.json()['id']).text_locations, '[]')

    def test_schema_upgrade_preserves_legacy_text_and_is_repeatable(self):
        self.material(); self.db.close()
        with self.engine.begin() as connection: connection.execute(text('ALTER TABLE materials DROP COLUMN text_locations'))
        upgrade_schema(self.engine); upgrade_schema(self.engine)
        with self.engine.connect() as connection:
            body, locations=connection.execute(text('SELECT text_content,text_locations FROM materials')).one()
            self.assertEqual((body, locations), (BODY,'[]'))


if __name__ == '__main__': unittest.main()
