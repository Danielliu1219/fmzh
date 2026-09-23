# -*- coding: utf-8 -*-
"""
《期末周别慌！》系统需求分析汇报 PPT 生成脚本
视觉：前端同一设计系统（纸 #f0ede3 × 墨 #0c1614 × 珊瑚 #e66348，墨线描边 + 硬阴影）
运行：python -X utf8 build_ppt.py
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, 'assets')
FIGS = os.path.join(HERE, '..', 'docs', '需求分析', 'figures')

# ---------- 设计系统 ----------
PAPER = RGBColor(0xF0, 0xED, 0xE3)
CARD = RGBColor(0xF8, 0xF5, 0xEC)
SOFT = RGBColor(0xE9, 0xE5, 0xD7)
LINE = RGBColor(0xD9, 0xD4, 0xC3)
INK = RGBColor(0x0C, 0x16, 0x14)
INK2 = RGBColor(0x45, 0x4E, 0x48)
INK3 = RGBColor(0x8B, 0x90, 0x84)
CORAL = RGBColor(0xE6, 0x63, 0x48)
CORAL_DEEP = RGBColor(0xC2, 0x47, 0x2F)
CORAL_SOFT = RGBColor(0xF8, 0xE1, 0xD7)
CORAL_BRIGHT = RGBColor(0xF0, 0x84, 0x6B)
GREEN = RGBColor(0x35, 0x55, 0x4C)
GREEN_SOFT = RGBColor(0xDF, 0xE8, 0xE2)
AMBER = RGBColor(0xA0, 0x6A, 0x22)
AMBER_SOFT = RGBColor(0xF2, 0xE7, 0xCF)
RED = RGBColor(0xB6, 0x3A, 0x2E)
RED_SOFT = RGBColor(0xF6, 0xDD, 0xD7)
TEAL = RGBColor(0x3F, 0x6F, 0x66)
TEAL_SOFT = RGBColor(0xDD, 0xEA, 0xE5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SERIF = 'STZhongsong'   # 华文中宋
SANS = '微软雅黑'
NUM = 'Bahnschrift'     # 数字/西文几何字体（近似 Space Grotesk）

PAGE_W, PAGE_H = 13.333, 7.5
MX, CW = 0.6, 13.333 - 1.2
SH = 0.055  # 硬阴影偏移

prs = Presentation()
prs.slide_width = Inches(PAGE_W)
prs.slide_height = Inches(PAGE_H)
BLANK = prs.slide_layouts[6]


def _set_run(run, size=11, bold=False, color=INK, ea=SANS, latin=NUM, spc=None):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = latin
    rPr = run._r.get_or_add_rPr()
    ea_el = rPr.find(qn('a:ea'))
    if ea_el is None:
        ea_el = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea_el)
    ea_el.set('typeface', ea)
    if spc:
        rPr.set('spc', str(int(spc * 100)))


def txt(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, wrap=True):
    """paras: list of dicts {runs:[(text,size,bold,color,ea,latin,spc)], align, ls, sa}"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get('align', PP_ALIGN.LEFT)
        if p.get('ls'):
            para.line_spacing = p['ls']
        if p.get('sa') is not None:
            para.space_after = Pt(p['sa'])
        for r in p['runs']:
            text, size, bold, color, ea, latin, spc = (list(r) + [None] * 7)[:7]
            run = para.add_run()
            run.text = text
            _set_run(run, size=size, bold=bold, color=color,
                     ea=ea or SANS, latin=latin or NUM, spc=spc)
    return tb


def box(slide, x, y, w, h, fill=CARD, line_color=INK, line_w=1.25,
        shadow=False, round_=False, radius=0.06, sh_color=None):
    if shadow:
        s = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
            Inches(x + SH), Inches(y + SH), Inches(w), Inches(h))
        if round_:
            try:
                s.adjustments[0] = radius
            except Exception:
                pass
        s.fill.solid()
        s.fill.fore_color.rgb = sh_color or INK
        s.line.fill.background()
        s.shadow.inherit = False
    sp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if round_:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    return sp


def chip(slide, x, y, w, h, text, fill=CORAL_SOFT, color=CORAL_DEEP, size=11,
         bold=True, line_color=None, round_=True, ea=SANS, latin=NUM, shadow=False):
    c = box(slide, x, y, w, h, fill=fill, line_color=line_color,
            shadow=shadow, round_=round_, radius=0.5 if round_ else 0.06)
    tf = c.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    _set_run(run, size=size, bold=bold, color=color, ea=ea, latin=latin)
    return c


def pic(slide, path, x, y, w=None, h=None):
    return slide.shapes.add_picture(path, Inches(x), Inches(y),
                                    Inches(w) if w else None, Inches(h) if h else None)


def fit_pic(path, box_w, box_h):
    """按比例放进给定框，返回 (w, h)"""
    with Image.open(path) as im:
        iw, ih = im.size
    ar = iw / ih
    w, h = box_w, box_w / ar
    if h > box_h:
        h, w = box_h, box_h * ar
    return w, h


def line(slide, x1, y1, x2, y2, color=INK, w=1.0):
    ln = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color
    ln.line.width = Pt(w)
    ln.shadow.inherit = False
    return ln


def header(slide, num, title):
    chip(slide, MX, 0.34, 0.46, 0.46, num, fill=CORAL, color=WHITE, size=16)
    txt(slide, 1.24, 0.36, 8.2, 0.5,
        [{'runs': [(title, 24, True, INK, SERIF)]}], anchor=MSO_ANCHOR.MIDDLE)
    line(slide, MX, 0.98, MX + CW, 0.98, INK, 1.25)


def footer(slide, page):
    txt(slide, 11.33, 7.13, 1.4, 0.3,
        [{'runs': [(page, 9, True, INK2, SANS, NUM)], 'align': PP_ALIGN.RIGHT}])


def new_slide(num, title):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = PAPER
    header(s, num, title)
    footer(s, num)
    return s


def browser(slide, x, y, w, h, img_path, caption, url_text='localhost:5173'):
    """墨线浏览器框：标题栏 + 截图 + 底部说明条"""
    frame = box(slide, x, y, w, h, fill=CARD, line_color=INK, line_w=1.25, shadow=True)
    # 标题栏
    box(slide, x + 1.25, y + 1.25, w - 2.5, 0.3, fill=SOFT, line_color=None)
    box(slide, x + 1.25, y, w - 2.5, 0.3, fill=SOFT, line_color=None)
    for i, c in enumerate([CORAL, AMBER, GREEN]):
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL,
                                     Inches(x + 0.12 + i * 0.17), Inches(y + 0.09),
                                     Inches(0.085), Inches(0.085))
        dot.fill.solid()
        dot.fill.fore_color.rgb = c
        dot.line.fill.background()
        dot.shadow.inherit = False
    txt(slide, x + 0.85, y + 0.035, w - 1.6, 0.24,
        [{'runs': [(url_text, 8, False, INK3, SANS, NUM)], 'align': PP_ALIGN.CENTER}],
        anchor=MSO_ANCHOR.MIDDLE)
    # 截图
    cap_h = 0.34
    area_w, area_h = w - 0.24, h - 0.46 - cap_h
    pw, ph = fit_pic(img_path, area_w, area_h)
    pic(slide, img_path, x + (w - pw) / 2, y + 0.42 + (area_h - ph) / 2, pw, ph)
    # 底部说明条
    txt(slide, x + 0.12, y + h - cap_h + 0.05, w - 0.24, cap_h - 0.1,
        [{'runs': [(caption, 9.5, True, INK2)], 'align': PP_ALIGN.CENTER}],
        anchor=MSO_ANCHOR.MIDDLE)
    return frame


# ============================================================ 1 封面
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = PAPER
line(s, MX, 0.62, MX + CW, 0.62, INK, 1.5)
box(s, MX, 0.44, 0.09, 0.36, fill=CORAL, line_color=None)
txt(s, MX, 0.28, 8, 0.3, [{'runs': [('四川大学计算机学院 · 2024 级', 10, True, INK2, SANS, NUM, 2)]}])
txt(s, 8.33, 0.28, 4.4, 0.3,
    [{'runs': [('REQUIREMENTS ANALYSIS · 2026-09', 10, True, INK3, SANS, NUM, 2)],
      'align': PP_ALIGN.RIGHT}])
# 右侧装饰：珊瑚色块 + 巨大感叹号
box(s, 9.45, 1.35, 3.0, 3.0, fill=CORAL, line_color=INK, line_w=1.5, shadow=True, round_=True, radius=0.05)
txt(s, 9.45, 1.7, 3.0, 2.4,
    [{'runs': [('！', 96, True, WHITE, SERIF)], 'align': PP_ALIGN.CENTER}])
txt(s, 9.45, 3.95, 3.0, 0.4,
    [{'runs': [('V1.2 · 2026-09', 11, True, CORAL_BRIGHT, SANS, NUM, 2)], 'align': PP_ALIGN.CENTER}])
# 水印
txt(s, 0.4, 5.9, 8, 1.4,
    [{'runs': [('FMZH', 110, True, RGBColor(0xE4, 0xE1, 0xD4), SANS, NUM)]}])
# 主标题块
chip(s, MX, 1.75, 3.3, 0.46, 'AI 智能备考网站 · 系统需求分析汇报',
     fill=CORAL_SOFT, color=CORAL_DEEP, size=12, shadow=True)
txt(s, MX, 2.45, 8.6, 1.6,
    [{'runs': [('期末周别慌', 64, True, INK, SERIF), ('！', 64, True, CORAL, SERIF)]}])
txt(s, MX, 3.75, 8.6, 0.5,
    [{'runs': [('面向大学生期末复习的一站式智能备考网站', 17, False, INK2)]}])
# 底部信息带
line(s, MX, 6.15, MX + CW, 6.15, LINE, 1.0)
txt(s, MX, 6.35, 12.13, 0.8, [
    {'runs': [('汇报人：刘添屹', 11.5, True, INK), ('    ·    小组成员：刘添屹  许丽媛  胡歆桐  张诗琪  宋跃月', 11, False, INK2)],
     'align': PP_ALIGN.LEFT},
    {'runs': [('指导老师：陈正茂    ·    《问题求解实战》课程', 11, False, INK2)], 'sa': 2},
])

# ============================================================ 2 目录
s = new_slide('01', '目录')
items = ['项目背景', '项目概述', '主要功能', '核心业务逻辑',
         '非功能需求', '原型展示', '进度与计划']
y = 1.38
for i, t in enumerate(items):
    chip(s, MX + 0.05, y, 0.52, 0.52, '0' + str(i + 1), fill=CORAL, color=WHITE, size=15)
    txt(s, 1.42, y, 5.0, 0.52,
        [{'runs': [(t, 20, True, INK, SERIF)]}], anchor=MSO_ANCHOR.MIDDLE)
    y += 0.78
# 右侧速览卡（几何为组长手调版：卡更窄更靠右，数字与文字紧贴）
box(s, 7.77, 1.51, 3.74, 4.6, fill=CARD, line_color=INK, line_w=1.25, shadow=True, round_=True)
chip(s, 8.02, 1.81, 2.6, 0.44, '数字速览', fill=INK, color=CORAL_BRIGHT, size=12)
stats = [('67', '条用户需求（连续编号）'), ('57', '个功能用例'),
         ('72', '条需求对照记录'), ('87', '页需求分析文档')]
y = 2.69
for n, l in stats:
    txt(s, 8.23, y, 1.5, 0.75, [{'runs': [(n, 30, True, CORAL, SANS, NUM)]}],
        anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 8.75, y, 3.4, 0.75,
        [{'runs': [(l, 11.5, True, INK2)], 'ls': 1.05}], anchor=MSO_ANCHOR.MIDDLE)
    y += 0.85

# ============================================================ 3 项目背景
s = new_slide('02', '项目背景')
txt(s, MX, 1.16, 12.13, 0.42,
    [{'runs': [('期末复习是大学生压力最集中的场景——', 14, True, INK),
               ('复习信息高度分散，学生难以在短时间内整合并制定有效计划。', 14, False, INK2)]}])
pains = [
    ('信息分散', '考表在教务处网站、重点散落在聊天记录、资料分布在多个渠道'),
    ('复习没重点', '历年题、笔记、PPT 多源异构，不知道从哪里开始'),
    ('计划靠感觉', '缺少科学规划，复习前松后紧、低效反复'),
    ('考前很焦虑', '任务积压、无从下手，越临近考试越慌乱'),
]
for i, (t, d) in enumerate(pains):
    x = MX + (i % 2) * 3.22
    y = 1.75 + (i // 2) * 1.72
    box(s, x, y, 3.06, 1.55, fill=CARD, line_color=INK, line_w=1.25, shadow=True, round_=True)
    chip(s, x + 0.18, y + 0.16, 1.0, 0.34, f'难题 0{i+1}', fill=RED_SOFT, color=RED, size=9.5)
    txt(s, x + 0.18, y + 0.62, 2.7, 0.8, [
        {'runs': [(t, 15, True, INK)], 'sa': 3},
        {'runs': [(d, 9.5, False, INK2)], 'ls': 1.12},
    ])
# 右侧：技术机遇
box(s, 7.25, 1.75, 5.48, 3.4, fill=GREEN, line_color=INK, line_w=1.25, shadow=True, round_=True)
chip(s, 7.55, 2.0, 1.9, 0.42, '技术机遇', fill=CORAL, color=WHITE, size=12)
txt(s, 7.55, 2.62, 4.9, 2.4, [
    {'runs': [('大模型技术日趋成熟', 16, True, WHITE)], 'sa': 6},
    {'runs': [('·  自然语言 → 结构化学习画像', 11, False, RGBColor(0xD9, 0xE4, 0xDE))], 'ls': 1.15, 'sa': 4},
    {'runs': [('·  多源资料 → 考点与权重分析', 11, False, RGBColor(0xD9, 0xE4, 0xDE))], 'ls': 1.15, 'sa': 4},
    {'runs': [('·  个性化生成 → 复习计划 / 自测题', 11, False, RGBColor(0xD9, 0xE4, 0xDE))], 'ls': 1.15, 'sa': 8},
    {'runs': [('DeepSeek 大模型 API 接入，成本可控、', 10.5, False, RGBColor(0xB9, 0xC9, 0xC1))], 'ls': 1.15},
    {'runs': [('内置 Mock 模式保证演示可靠', 10.5, False, RGBColor(0xB9, 0xC9, 0xC1))], 'ls': 1.15},
])
# 底部方案带
box(s, MX, 5.5, 12.13, 1.28, fill=CORAL_SOFT, line_color=INK, line_w=1.25, shadow=True, round_=True)
txt(s, 0.85, 5.66, 11.6, 1.0, [
    {'runs': [('我们的方案：', 14, True, CORAL_DEEP),
              ('打造一站式智能备考网站 —— 上传复习资料 + 描述学习状态，', 13.5, True, INK)], 'sa': 3},
    {'runs': [('AI 生成个性化复习重点、按天计划与自测题，每天打卡冲刺期末。', 13.5, True, INK)]},
])

# ============================================================ 4 项目概述
s = new_slide('03', '项目概述')
box(s, MX, 1.2, 12.13, 1.42, fill=INK, line_color=None, shadow=True, round_=True)
txt(s, 0.9, 1.36, 11.6, 1.1, [
    {'runs': [('「期末周别慌！」', 15, True, CORAL_BRIGHT, SERIF),
              ('面向大学生期末复习场景：学生上传多源复习资料、描述备考目标与学习状态，', 12.5, False, WHITE)], 'ls': 1.2},
    {'runs': [('系统调用 DeepSeek 大模型生成个性化复习重点、按天计划与自测题，配套打卡与统计，一站式科学备考。', 12.5, False, WHITE)], 'ls': 1.2},
])
txt(s, MX, 2.82, 6, 0.4, [{'runs': [('完整复习闭环', 14, True, INK, SERIF)]}])
steps = ['录入考试', '上传资料', '描述状态', 'AI 提取画像', '生成分析', '每日打卡', '统计反馈']
x = MX
for i, st in enumerate(steps):
    box(s, x, 3.3, 1.52, 0.92, fill=CARD, line_color=INK, line_w=1.25, shadow=True, round_=True, radius=0.12)
    txt(s, x, 3.42, 1.52, 0.3, [{'runs': [(str(i + 1), 11, True, CORAL, SANS, NUM)],
                                'align': PP_ALIGN.CENTER}])
    txt(s, x, 3.72, 1.52, 0.4, [{'runs': [(st, 11.5, True, INK)], 'align': PP_ALIGN.CENTER}])
    if i < 6:
        txt(s, x + 1.5, 3.5, 0.24, 0.4, [{'runs': [('→', 14, True, CORAL, SANS, NUM)],
                                          'align': PP_ALIGN.CENTER}])
    x += 1.76
txt(s, MX, 4.36, 12.13, 0.34,
    [{'runs': [('↺  打卡进度与剩余天数变化 → 动态调整计划与权重，形成闭环', 11.5, True, INK2)],
      'align': PP_ALIGN.CENTER}])
line(s, MX, 4.9, MX + CW, 4.9, LINE, 1.0)
txt(s, MX, 5.02, 3, 0.4, [{'runs': [('系统角色', 14, True, INK, SERIF)]}])
roles = [
    ('学生', '主要使用者：考试 / 资料 / AI 分析 / 打卡 / 特色功能'),
    ('游客', '仅可注册与登录，注册后成为学生'),
    ('管理员', '开发小组担任：发布统一考表、维护科目、治理违规内容'),
]
x = MX
for i, (t, d) in enumerate(roles):
    box(s, x, 5.5, 3.7, 1.05, fill=CARD, line_color=INK, line_w=1.25, shadow=True, round_=True)
    chip(s, x + 0.18, 5.66, 0.92, 0.34, t, fill=CORAL_SOFT, color=CORAL_DEEP, size=10.5)
    txt(s, x + 1.24, 5.62, 2.34, 0.85, [{'runs': [(d, 8.5, False, INK2)], 'ls': 1.1}])
    x += 3.9
box(s, 12.35, 5.5, 0.42, 1.05, fill=None, line_color=None)
txt(s, 10.55, 5.62, 2.15, 0.85, [
    {'runs': [('外部系统', 10, True, INK2, SANS, NUM, 1)], 'sa': 3},
    {'runs': [('DeepSeek API', 8.5, False, INK2)], 'ls': 1.1},
    {'runs': [('教务处网站', 8.5, False, INK2)], 'ls': 1.1},
])

# ============================================================ 5 创新点
s = new_slide('04', '创新点')
inno = [
    ('两级权重体系', '「资料权重 × 知识点权重」分离——可信度与重要性各算各的，分析结果既有依据又有个性'),
    ('画像驱动的差异化方案', '勾选表单 + 自然语言快速建画像，不同目标与基础得到不同复习策略'),
    ('历年题驱动的考点分析', '大模型提取历年题必考点与常考题型，复习重点与自测题贴合真实考试'),
]
for i, (t, d) in enumerate(inno):
    x = MX + i * 4.17
    box(s, x, 1.55, 3.95, 2.15, fill=CARD, line_color=INK, line_w=1.25, shadow=True, round_=True)
    txt(s, x + 0.22, 1.72, 3.5, 0.6, [{'runs': [('0' + str(i + 1), 24, True, CORAL, SANS, NUM),
                                               ('  ' + t, 15, True, INK, SERIF)]}])
    txt(s, x + 0.22, 2.42, 3.55, 1.15, [{'runs': [(d, 10, False, INK2)], 'ls': 1.2}])
inno2 = [
    ('人文关怀功能', '压力指数、临时救急模式、今日最重要的三件事、情绪记录，缓解考前焦虑'),
    ('教务处考表一键导入', '内嵌 WebView 注入脚本读取统一考表，免除手动录入负担'),
]
for i, (t, d) in enumerate(inno2):
    x = MX + i * 6.22
    box(s, x, 4.05, 5.95, 1.55, fill=CARD, line_color=INK, line_w=1.25, shadow=True, round_=True)
    txt(s, x + 0.25, 4.22, 5.4, 0.6, [{'runs': [('0' + str(i + 4), 24, True, CORAL, SANS, NUM),
                                               ('  ' + t, 15, True, INK, SERIF)]}])
    txt(s, x + 0.25, 4.9, 5.45, 0.6, [{'runs': [(d, 10, False, INK2)], 'ls': 1.2}])
box(s, MX, 5.95, 12.13, 0.95, fill=TEAL_SOFT, line_color=None, round_=True)
txt(s, 0.9, 6.14, 11.6, 0.6, [
    {'runs': [('定位差异：', 12.5, True, TEAL),
              ('不做通用日程管理，只做「期末冲刺」这一个场景——与超级课程表、滴答清单等形成差异。', 12.5, False, INK)]}])

# ============================================================ 6 功能总览
s = new_slide('05', '主要功能')
mods = [
    ('0', '公共功能', '5 用例', '鉴权 / 数据隔离 / 二次确认 / 日志 / 上传校验'),
    ('1', '账号管理', '5 用例', '注册 / 登录 / 锁定找回 / 资料 / 注销'),
    ('2', '考试与日程管理', '13 用例', '录入考试 / 任务 / 日历 / 番茄钟 / ICS'),
    ('3', '复习资料管理', '6 用例', '上传 / 标注来源 / 分类 / 搜索 / 错题本'),
    ('4', 'AI 智能分析', '14 用例', '画像 / 权重 / 提纲 / 计划 / 自测题'),
    ('5', '学习统计', '3 用例', '打卡 / 统计图表 / 备考周报'),
    ('6', '提醒通知', '2 用例', '考前与任务提醒 / 通知管理'),
    ('7', '特色功能', '6 用例', '压力指数 / 救急模式 / 情绪记录'),
    ('8', '管理端功能', '3 用例', '发布考表 / 维护科目 / 内容治理'),
]
y = 1.22
for n, t, c, _d in mods:
    if t == 'AI 智能分析':
        box(s, MX, y, 5.3, 0.52, fill=INK, line_color=INK, shadow=True, round_=True, radius=0.5)
        txt(s, MX + 0.18, y, 0.5, 0.52, [{'runs': [(n, 13, True, CORAL_BRIGHT, SANS, NUM)]}],
            anchor=MSO_ANCHOR.MIDDLE)
        txt(s, MX + 0.75, y, 2.85, 0.52, [{'runs': [(t, 12.5, True, WHITE)]}],
            anchor=MSO_ANCHOR.MIDDLE)
        txt(s, MX + 3.05, y, 2.1, 0.52, [{'runs': [(c, 10.5, True, CORAL_BRIGHT, SANS, NUM)],
                                          'align': PP_ALIGN.RIGHT}], anchor=MSO_ANCHOR.MIDDLE)
    else:
        box(s, MX, y, 5.3, 0.52, fill=CARD, line_color=INK, line_w=1.0, shadow=False, round_=True, radius=0.5)
        txt(s, MX + 0.18, y, 0.5, 0.52, [{'runs': [(n, 13, True, CORAL, SANS, NUM)]}],
            anchor=MSO_ANCHOR.MIDDLE)
        txt(s, MX + 0.75, y, 2.85, 0.52, [{'runs': [(t, 12.5, True, INK)]}],
            anchor=MSO_ANCHOR.MIDDLE)
        txt(s, MX + 3.05, y, 2.1, 0.52, [{'runs': [(c, 10.5, True, INK3, SANS, NUM)],
                                          'align': PP_ALIGN.RIGHT}], anchor=MSO_ANCHOR.MIDDLE)
    y += 0.585
# 右：顶层用例图
fig = os.path.join(FIGS, 'fig_3_1_1.png')
with Image.open(fig) as im:
    fiw, fih = im.size
fig_h = 5.25
fig_w = fig_h * fiw / fih
pic(s, fig, 12.73 - fig_w, 1.28, fig_w, fig_h)
chip(s, 7.0, 1.18, 1.6, 0.34, '系统顶层用例图', fill=SOFT, color=INK2, size=9.5)
stats = [('67', '条用户需求'), ('57', '个功能用例'), ('72', '条对照记录'), ('9', '大功能模块')]
for i, (n, l) in enumerate(stats):
    x = MX + i * 3.15
    box(s, x, 6.62, 2.95, 0.5, fill=INK, line_color=None, round_=True, radius=0.5)
    txt(s, x + 0.08, 6.62, 2.79, 0.5,
        [{'runs': [(n + '  ', 17, True, CORAL_BRIGHT, SANS, NUM),
                   (l, 10.5, True, WHITE)], 'align': PP_ALIGN.CENTER}],
        anchor=MSO_ANCHOR.MIDDLE)

# ============================================================ 7 AI 智能分析（核心）
s = new_slide('06', '核心模块 · AI 智能分析')
txt(s, MX, 1.14, 12.13, 0.38, [{'runs': [('从资料到个性化方案：五步链路', 13.5, True, INK, SERIF)]}])
chain = ['上传资料', '解析文本', '权重计算', '方案生成', '落为任务']
x = MX
for i, st in enumerate(chain):
    box(s, x, 1.62, 2.1, 0.78, fill=CARD, line_color=INK, line_w=1.25, shadow=True, round_=True, radius=0.14)
    txt(s, x, 1.74, 2.1, 0.3, [{'runs': [(str(i + 1), 10.5, True, CORAL, SANS, NUM)],
                               'align': PP_ALIGN.CENTER}])
    txt(s, x, 2.0, 2.1, 0.3, [{'runs': [(st, 11.5, True, INK)], 'align': PP_ALIGN.CENTER}])
    if i < 4:
        txt(s, x + 2.06, 1.84, 0.3, 0.35, [{'runs': [('→', 13, True, CORAL, SANS, NUM)],
                                            'align': PP_ALIGN.CENTER}])
    x += 2.4
ctrls = ['固定提示词', '固定 JSON 输出', '失败自动重试']
cx = MX
for c in ctrls:
    chip(s, cx, 2.56, 1.35, 0.34, c, fill=SOFT, color=INK2, size=9.5, shadow=False)
    cx += 1.52
# 左：两级权重
txt(s, MX, 3.12, 6, 0.36, [{'runs': [('两级权重体系', 13.5, True, INK, SERIF)]}])
box(s, MX, 3.5, 6.0, 1.66, fill=CARD, line_color=INK, line_w=1.25, shadow=True, round_=True)
chip(s, MX + 0.2, 3.66, 2.1, 0.38, '资料权重 · 可不可信', fill=CORAL_SOFT, color=CORAL_DEEP, size=10.5)
txt(s, MX + 0.2, 4.18, 5.6, 0.85, [
    {'runs': [('老师标注重点  ＞  历年试题 / 练习题', 11.5, True, INK)], 'ls': 1.25, 'sa': 3},
    {'runs': [('＞  平时作业 / 课堂笔记  ＞  教材章节 / PPT / 实验报告', 11.5, True, INK2)], 'ls': 1.25},
])
box(s, MX, 5.28, 6.0, 1.44, fill=CARD, line_color=INK, line_w=1.25, shadow=True, round_=True)
chip(s, MX + 0.2, 5.44, 2.45, 0.38, '知识点权重 · 重不重要', fill=GREEN_SOFT, color=GREEN, size=10.5)
txt(s, MX + 0.2, 5.96, 5.6, 0.6,
    [{'runs': [('历年题必考点 → 基础值 → 用户画像调整（目标 · 薄弱章节）', 11.5, True, INK)], 'ls': 1.2}])
# 右：三档目标 + 输出
txt(s, 6.85, 3.12, 6, 0.36, [{'runs': [('分级推荐策略 · 三档备考目标', 13.5, True, INK, SERIF)]}])
tiers = [
    ('保及格', GREEN, GREEN_SOFT, '基础薄弱者优先：高频基础考点，拿稳基本分'),
    ('稳中等', AMBER, AMBER_SOFT, '常规考点全覆盖，按权重均衡安排'),
    ('冲高分', CORAL_DEEP, CORAL_SOFT, '综合题与难题专项，追求高覆盖率'),
]
y = 3.5
for t, c, cf, d in tiers:
    box(s, 6.85, y, 5.88, 0.82, fill=cf, line_color=INK, line_w=1.0, round_=True, radius=0.5)
    txt(s, 7.05, y, 1.35, 0.82, [{'runs': [(t, 13.5, True, c, SERIF)]}], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 8.35, y + 0.08, 4.3, 0.68, [{'runs': [(d, 9.5, False, INK2)], 'ls': 1.12}],
        anchor=MSO_ANCHOR.MIDDLE)
    y += 0.94
box(s, 6.85, 6.36, 5.88, 0.5, fill=INK, line_color=None, round_=True, radius=0.5)
txt(s, 7.05, 6.36, 5.6, 0.5,
    [{'runs': [('输出：复习提纲 · 知识点权重 · 按天计划 · 自测题 —— 每个结论带依据', 10.5, True, CORAL_BRIGHT, SANS, NUM)]}],
    anchor=MSO_ANCHOR.MIDDLE)

# ============================================================ 8 特色功能与公式
s = new_slide('07', '特色功能与核心公式')
# 左：压力指数
txt(s, MX, 1.14, 6.5, 0.38, [{'runs': [('期末压力指数（0~100）', 14.5, True, INK, SERIF)]}])
box(s, MX, 1.56, 6.55, 2.9, fill=INK, line_color=INK, shadow=True, round_=True)
txt(s, MX + 0.25, 1.74, 6.0, 1.1, [
    {'runs': [('压力指数 = ', 14, True, WHITE),
              ('30×密集度 + 30×积压率 + 20×(1−完成率) + 20×临近度', 14, True, CORAL_BRIGHT, SANS, NUM)], 'ls': 1.15},
])
subs = [
    ('考试密集度', 'min(1, 未来 7 天考试数 ÷ 4)'),
    ('任务积压率', '逾期未完成任务数 ÷ 总任务数'),
    ('任务完成率', '已完成任务数 ÷ 总任务数'),
    ('考试临近度', 'min(1, 21 ÷ 最近考试剩余天数)'),
]
for i, (t, f) in enumerate(subs):
    x = MX + 0.22 + (i % 2) * 3.2
    y = 2.52 + (i // 2) * 0.88
    txt(s, x, y, 3.0, 0.8, [
        {'runs': [(t + '  ', 10.5, True, CORAL_BRIGHT)], 'ls': 1.1},
        {'runs': [(f, 10.5, False, WHITE, SANS, NUM)], 'ls': 1.1},
    ])
txt(s, MX + 0.22, 4.28, 6.1, 0.3,
    [{'runs': [('边界约定：无考试取 0 · 考试当天取 1 · 无任务取 0', 8.5, False, RGBColor(0xB9, 0xC9, 0xC1))]}])
box(s, MX, 4.95, 6.55, 1.02, fill=CARD, line_color=INK, line_w=1.25, shadow=True, round_=True)
txt(s, MX + 0.2, 5.09, 6.15, 0.75, [
    {'runs': [('每日定时计算，按档提示：', 10.5, True, INK),
              ('<40 轻松 · 40~60 适中 · 60~80 偏压 · ≥80 高压', 10.5, True, INK2)], 'ls': 1.15, 'sa': 3},
    {'runs': [('≥70 建议开启救急模式', 9.5, True, CORAL_DEEP)], 'ls': 1.1},
])
# 右：情绪公式 + 特色列表
txt(s, 7.4, 1.14, 5.3, 0.38, [{'runs': [('情绪综合评分（0~5）', 14.5, True, INK, SERIF)]}])
box(s, 7.4, 1.56, 5.33, 2.06, fill=CARD, line_color=INK, line_w=1.25, shadow=True, round_=True)
txt(s, 7.62, 1.72, 4.9, 0.85, [
    {'runs': [('0.4×情绪项 + 0.3×完成项 + 0.2×压力项 + 0.1×达标项', 12.5, True, INK, SANS, NUM)], 'ls': 1.2, 'sa': 4},
    {'runs': [('四个分项均归一化到 0~5，总分数学上精确落在 0~5', 9.5, True, GREEN)], 'ls': 1.1},
])
txt(s, 7.62, 2.68, 4.9, 0.85, [
    {'runs': [('情绪项 = 5×(自评−1)÷4    完成项 = 5×当日完成率', 9.5, False, INK2)], 'ls': 1.2, 'sa': 2},
    {'runs': [('压力项 = 5−压力指数÷20   达标项 = 5×min(连续达标天数,7)÷7', 9.5, False, INK2)], 'ls': 1.2},
])
txt(s, 7.4, 3.86, 5.3, 0.38, [{'runs': [('其他特色功能', 13, True, INK, SERIF)]}])
feats = ['临时救急模式：时间紧张时聚焦高频基础考点', '今日最重要的三件事：权重 × 截止时间 × 压力指数',
         '考前重点回顾列表 · 情绪记录与节奏建议', '学习时段推荐：按画像分配每日复习时段']
y = 4.26
for f in feats:
    box(s, 7.4, y, 5.33, 0.5, fill=CARD, line_color=INK, line_w=1.0, round_=True, radius=0.5)
    txt(s, 7.58, y, 5.0, 0.5, [{'runs': [(f, 10, False, INK2)]}], anchor=MSO_ANCHOR.MIDDLE)
    y += 0.6

# ============================================================ 9 非功能需求
s = new_slide('08', '非功能需求')
cols = [
    ('性能', CORAL, [
        '数据库操作响应 ≤ 2 秒',
        '页面加载 ≤ 4 秒',
        '20MB 文件上传受理 ≤ 10 秒，解析后台异步完成',
        'AI 分析 10~60 秒异步执行，不阻塞用户',
        '支持 50 名用户并发使用',
        '压力指数、周报每日凌晨定时计算',
    ]),
    ('安全', GREEN, [
        '密码 bcrypt 哈希存储，不明文传输',
        'JWT 会话鉴权，有效期 7 天',
        '全站 HTTPS；API Key 仅存后端环境变量',
        '所有查询强制 user_id 过滤（数据隔离）',
        '连续 3 次密码错误锁定账号',
        '删除 / 注销二次确认 + 关键操作日志',
    ]),
    ('可靠与易用', TEAL, [
        '大模型调用超时 60 秒，失败自动重试 1 次',
        '服务不可用时降级提示，Mock 模式保演示',
        '数据库与上传文件每日定时备份',
        '现代浏览器全兼容，移动端响应式',
        '编辑杂志风设计语言，操作路径短',
        'AI 内容固定展示「仅供复习参考」提示',
    ]),
]
for i, (t, c, items) in enumerate(cols):
    x = MX + i * 4.17
    box(s, x, 1.32, 3.95, 5.05, fill=CARD, line_color=INK, line_w=1.25, shadow=True, round_=True)
    chip(s, x + 0.2, 1.5, 1.4, 0.42, t, fill=c, color=WHITE, size=12)
    txt(s, x + 0.2, 2.16, 3.55, 4.0,
        [{'runs': [('·  ' + it, 10.5, False, INK2)], 'ls': 1.18, 'sa': 8} for it in items])

# ============================================================ 10 数据描述
s = new_slide('09', '数据描述')
tables = [
    ('users', '用户', '账号（唯一）· bcrypt 密码哈希 · 昵称 / 学号 / 专业 / 年级'),
    ('exams', '考试安排', '科目 · 考试日期 · 地点 · 时长 · 备注'),
    ('materials', '复习资料', '标题 · 来源类型 · 文件与解析文本 · 标签'),
    ('user_profiles', '学习画像', '备考目标 · 学习基础 · 剩余天数 · 薄弱章节(JSON)'),
    ('analyses', '分析结果', '固定 JSON：权重 / 提纲 / 按天计划 / 自测题'),
    ('tasks', '学习任务', '所属考试 · 计划日期 · 时长 · 优先级 · 完成状态'),
]
for i, (n, cn, d) in enumerate(tables):
    x = MX + (i % 3) * 4.17
    y = 1.5 + (i // 3) * 2.1
    box(s, x, y, 3.95, 1.88, fill=CARD, line_color=INK, line_w=1.25, shadow=True, round_=True)
    txt(s, x + 0.2, y + 0.14, 3.55, 0.4,
        [{'runs': [(n, 14, True, INK, SANS, NUM), ('   ' + cn, 12.5, True, CORAL_DEEP)]}])
    txt(s, x + 0.2, y + 0.62, 3.6, 1.1, [{'runs': [(d, 9.5, False, INK2)], 'ls': 1.2}])
box(s, MX, 5.78, 12.13, 0.78, fill=INK, line_color=None, shadow=True, round_=True)
txt(s, 0.9, 5.78, 11.6, 0.78, [
    {'runs': [('数据约束：', 11.5, True, CORAL_BRIGHT),
              ('账号 / 学号唯一 · 增删改走事务 · 外键约束明确 · JSON Schema + Pydantic 双重校验 · 浮点两位', 11, False, WHITE)], 'ls': 1.1}],
    anchor=MSO_ANCHOR.MIDDLE)

# ============================================================ 11 需求分析成果
s = new_slide('10', '需求分析成果')
big = [
    ('67', '条用户需求', 'FMZH-USR-0001 ~ 0067 连续编号'),
    ('57', '个功能用例', 'FMZH-SRS-0.1.0 ~ 8.3.0'),
    ('72', '条对照记录', '67 条用户需求 + 5 条隐藏需求'),
    ('87', '页 · V1.2', '9 章 · 78 张表 · 68 张 UML 图'),
]
for i, (n, l, d) in enumerate(big):
    x = MX + i * 3.15
    box(s, x, 1.55, 2.95, 2.5, fill=CARD, line_color=INK, line_w=1.25, shadow=True, round_=True)
    txt(s, x, 1.85, 2.95, 0.8, [{'runs': [(n, 36, True, CORAL, SANS, NUM)], 'align': PP_ALIGN.CENTER}])
    txt(s, x, 2.75, 2.95, 0.4, [{'runs': [(l, 13.5, True, INK, SERIF)], 'align': PP_ALIGN.CENTER}])
    txt(s, x + 0.15, 3.2, 2.65, 0.7, [{'runs': [(d, 9, False, INK2)], 'align': PP_ALIGN.CENTER, 'ls': 1.15}])
box(s, MX, 4.45, 6.6, 2.3, fill=CARD, line_color=INK, line_w=1.25, shadow=True, round_=True)
txt(s, 0.85, 4.66, 11.6, 0.4, [{'runs': [('需求分析工作方法', 14, True, INK, SERIF)]}])
txt(s, 0.85, 5.18, 6.3, 1.35, [
    {'runs': [('用户需求列表（67 条）', 11.5, True, CORAL_DEEP),
              ('  →  系统需求分析（57 用例 + 活动图 + 数据字典）', 11.5, False, INK)], 'ls': 1.25, 'sa': 5},
    {'runs': [('  →  需求对照表（72 条，双向可追溯）', 11.5, False, INK)], 'ls': 1.25, 'sa': 6},
    {'runs': [('另补 5 条隐藏需求：数据隔离、操作日志等', 10.5, False, INK2)], 'ls': 1.2},
])
box(s, 6.85, 4.45, 5.88, 2.3, fill=CORAL_SOFT, line_color=INK, line_w=1.0, round_=True)
txt(s, 7.05, 4.66, 5.5, 0.4, [{'runs': [('文档已交付', 14, True, CORAL_DEEP, SERIF)]}])
txt(s, 7.05, 5.18, 5.5, 1.4, [
    {'runs': [('系统需求分析 docx', 12, True, INK)], 'ls': 1.2, 'sa': 3},
    {'runs': [('PDF 预览版', 12, True, INK)], 'ls': 1.2, 'sa': 3},
    {'runs': [('简要用户需求列表', 12, True, INK)], 'ls': 1.2},
])

# ============================================================ 12 原型展示 ①
s = new_slide('11', '原型展示 ①')
txt(s, MX, 1.14, 12.13, 0.38, [{'runs': [('已实现的系统界面 · 编辑杂志风', 13.5, True, INK, SERIF)]}])
browser(s, MX, 1.58, 5.95, 5.0, os.path.join(ASSETS, '01_login.png'),
        '登录 / 注册 · JWT 认证', 'localhost:5173/login')
browser(s, 6.78, 1.58, 5.95, 5.0, os.path.join(ASSETS, '02_home.png'),
        '首页聚合面板 · 今日任务 / 倒计时 / 进度 / 打卡（演示数据）', 'localhost:5173')

# ============================================================ 13 原型展示 ②
s = new_slide('12', '原型展示 ②')
shots = [
    ('03_exams.png', '考试管理 · 录入 / 进度 / 倒计时', 'localhost:5173/exams'),
    ('04_materials.png', '复习资料 · 上传 / 来源标注 / 解析', 'localhost:5173/materials'),
    ('05_analysis.png', 'AI 分析 · 知识点权重', 'localhost:5173/analysis'),
    ('05_analysis_scroll.png', 'AI 分析 · 按天计划 / 自测题', 'localhost:5173/analysis'),
]
for i, (fn, cap, url) in enumerate(shots):
    x = MX + (i % 2) * 6.2
    y = 1.48 + (i // 2) * 2.68
    browser(s, x, y, 5.95, 2.5 + (i // 2) * 0.42, os.path.join(ASSETS, fn), cap, url)

# ============================================================ 14 技术架构
s = new_slide('13', '技术架构')
stack = [
    ('浏览器', 'Chrome / Edge / Firefox，桌面端与移动端均可访问', CARD),
    ('前端', 'Vue 3 + Vite + Pinia + Vue Router + Element Plus + ECharts + Axios', CARD),
    ('后端', 'Python 3.10+ · FastAPI + SQLAlchemy 2 + Pydantic v2 · JWT 鉴权', CARD),
    ('数据库', 'SQLite（开发演示）→ MySQL 8.0（目标），切换仅改连接配置', GREEN_SOFT),
    ('外部服务', 'DeepSeek API（OpenAI 兼容 · Key 仅存后端环境变量 · 内置 Mock）', CORAL_SOFT),
    ('', '四川大学教务处网站（统一考表导入，WebView + 脚本注入）', CARD),
]
y = 1.24
for i, (t, d, fill) in enumerate(stack):
    box(s, MX, y, 5.5, 0.72, fill=fill, line_color=INK, line_w=1.0, round_=True, radius=0.5)
    txt(s, MX + 0.2, y, 1.35, 0.72, [{'runs': [(t, 12, True, INK, SERIF)]}], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, MX + 1.6, y + 0.05, 3.75, 0.62, [{'runs': [(d, 8.5, False, INK2)], 'ls': 1.05}],
        anchor=MSO_ANCHOR.MIDDLE)
    if i < 5:
        txt(s, MX + 0.2, y + 0.68, 0.3, 0.25, [{'runs': [('↓', 11, True, CORAL, SANS, NUM)]}])
    y += 0.94
fig = os.path.join(FIGS, 'fig_2_2_1.png')
with Image.open(fig) as im:
    fiw, fih = im.size
fig_h = 5.5
fig_w = fig_h * fiw / fih
pic(s, fig, 12.73 - fig_w, 1.24, fig_w, fig_h)
txt(s, 6.6, 6.82, 6.13, 0.3,
    [{'runs': [('系统总体构架图', 9, False, INK2)], 'align': PP_ALIGN.RIGHT}])
txt(s, MX, 6.82, 5.5, 0.3,
    [{'runs': [('前后端分离 · 接口文档自动生成 · 双 Mock 保演示', 9.5, False, INK2)]}])

# ============================================================ 15 分工与进度
s = new_slide('14', '进度与小组分工')
txt(s, MX, 1.14, 12.13, 0.38, [{'runs': [('当前进度', 14, True, INK, SERIF)]}])
dones = [
    ('需求分析', 'V1.2 已交付：87 页 · 9 章 · 78 表 · 68 图'),
    ('MVP 主线', '注册登录 → 录考试 → 传资料 → AI 分析 → 打卡 → 统计，全流程打通'),
    ('双 Mock 设计', '前端假后端 / 后端假大模型，两开关正交，无网络无 Key 也能演示'),
]
y = 1.58
for t, d in dones:
    box(s, MX, y, 4.9, 0.92, fill=GREEN_SOFT, line_color=INK, line_w=1.0, round_=True, radius=0.5)
    txt(s, MX + 0.2, y, 4.5, 0.92, [
        {'runs': [('✓ ' + t + '  ', 11.5, True, GREEN), (d, 9.5, False, INK2)], 'ls': 1.1}],
        anchor=MSO_ANCHOR.MIDDLE)
    y += 1.04
txt(s, MX, 4.78, 5, 0.38, [{'runs': [('待开发（已分配至组员）', 13, True, INK, SERIF)]}])
box(s, MX, 5.22, 4.9, 1.5, fill=CARD, line_color=INK, line_w=1.0, round_=True)
txt(s, MX + 0.2, 5.36, 4.5, 1.25, [
    {'runs': [('提醒通知 · 特色功能 · 管理端三个模块；', 10, False, INK2)], 'ls': 1.25, 'sa': 3},
    {'runs': [('资料类型补齐（教材章节 / 实验报告）；', 10, False, INK2)], 'ls': 1.25, 'sa': 3},
    {'runs': [('任务编辑、考表导入等扩展项，全部录入 README 状态表。', 10, False, INK2)], 'ls': 1.25},
])
# 右：分工表
txt(s, 6.4, 1.14, 6.3, 0.38, [{'runs': [('小组分工', 14, True, INK, SERIF)]}])
team = [
    ('刘添屹', '组长', '现有成果全部完成；后续架构把关、代码审查、全文统稿'),
    ('许丽媛', '组员', '账号管理完善 + 考试与日程扩展'),
    ('胡歆桐', '组员', '资料管理扩展 + AI 分析扩展'),
    ('张诗琪', '组员', '学习统计 + 提醒通知 + 番茄钟 + 自测题判分'),
    ('宋跃月', '组员', '特色功能 + 管理端 + 二次确认/日志 + 联调测试'),
]
y = 1.58
for name, role, d in team:
    box(s, 6.4, y, 6.33, 0.78, fill=CARD, line_color=INK, line_w=1.0, round_=True, radius=0.5)
    txt(s, 6.62, y, 1.15, 0.78, [{'runs': [(name, 12.5, True, INK, SERIF)],
                                  'align': PP_ALIGN.LEFT}], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 7.75, y, 0.7, 0.78, [{'runs': [(role, 9, True, CORAL_DEEP)],
                                 'align': PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 8.55, y, 4.05, 0.78, [{'runs': [(d, 9.5, False, INK2)], 'ls': 1.05}],
        anchor=MSO_ANCHOR.MIDDLE)
    y += 0.885

# ============================================================ 16 后续计划与结束
s = new_slide('15', '后续计划')
plans = [
    ('① 前后端联调', 'USE_MOCK 切回 false，跑通注册 → 录入 → 上传 → AI 分析 → 打卡全流程'),
    ('② 组员开发', '按分工表认领任务，Git 分支协作 + Pull Request 合并'),
    ('③ 公网部署', 'Nginx + Uvicorn 部署到公网，内网可演示兜底'),
    ('④ 文档回归', '任何公式 / 口径 / 接口改动，同步更新需求文档与交接文档'),
]
for i, (t, d) in enumerate(plans):
    x = MX + i * 3.15
    box(s, x, 1.55, 2.95, 2.3, fill=CARD, line_color=INK, line_w=1.25, shadow=True, round_=True)
    txt(s, x + 0.2, 1.8, 2.55, 0.5, [{'runs': [(t, 13.5, True, CORAL_DEEP, SERIF)]}])
    txt(s, x + 0.2, 2.35, 2.6, 1.35, [{'runs': [(d, 9.5, False, INK2)], 'ls': 1.25}])
box(s, MX, 4.35, 12.13, 2.35, fill=INK, line_color=INK, shadow=True, round_=True)
txt(s, 0.9, 4.75, 11.6, 0.9,
    [{'runs': [('谢谢聆听', 40, True, WHITE, SERIF)], 'align': PP_ALIGN.CENTER}])
txt(s, 0.9, 5.75, 11.6, 0.5,
    [{'runs': [('恳请老师与同学们批评指正', 15, False, CORAL_BRIGHT)], 'align': PP_ALIGN.CENTER}])

out = os.path.join(HERE, '期末周别慌_需求分析汇报.pptx')
prs.save(out)
print('SAVED:', out, '| slides:', len(prs.slides.__iter__.__self__._sldIdLst))
