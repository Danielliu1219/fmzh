# -*- coding: utf-8 -*-
"""生成《期末周别慌！》系统需求分析.docx（图表缺图时以占位文字替代，渲染后重跑即可嵌入）"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from content import *  # noqa: F401,F403

HERE = os.path.dirname(os.path.abspath(__file__))
FIG_DIR = os.path.join(HERE, 'figures')
OUT = os.path.join(HERE, '..', '..', '期末周别慌_系统需求分析.docx')

CENTER = WD_ALIGN_PARAGRAPH.CENTER
LEFT = WD_ALIGN_PARAGRAPH.LEFT


def set_font(run, cn='宋体', en='Times New Roman', size=12, bold=False, color=None):
    run.font.name = en
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn('w:eastAsia'), cn)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def setup_styles(doc):
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
    normal.font.size = Pt(12)
    for name, sz in [('Heading 1', 16), ('Heading 2', 14), ('Heading 3', 12), ('Heading 4', 12)]:
        st = doc.styles[name]
        st.font.name = 'Times New Roman'
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '黑体')
        st.font.size = Pt(sz)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)


def para(doc, text='', cn='宋体', size=12, bold=False, align=LEFT, indent=True,
         space_after=4, space_before=0, color=None, line=1.25):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = line
    if indent and text and align == LEFT:
        pf.first_line_indent = Pt(size * 2)
    r = p.add_run(text)
    set_font(r, cn=cn, size=size, bold=bold, color=color)
    return p


def _heading(doc, text, level, size):
    p = doc.add_heading(text, level=level)
    pf = p.paragraph_format
    pf.space_before = Pt(12 if level == 1 else (8 if level == 2 else 6))
    pf.space_after = Pt(8 if level == 1 else (6 if level == 2 else 3))
    pf.keep_with_next = True  # 标题不与正文分离，避免页尾孤行留白
    # run 级显式字体：不依赖样式表，Word/WPS 均按此显示
    for r in p.runs:
        set_font(r, cn='黑体', size=size, bold=True, color='000000')
    return p


def h1(doc, text):
    return _heading(doc, text, 1, 16)


def h2(doc, text):
    return _heading(doc, text, 2, 14)


def h3(doc, text):
    return _heading(doc, text, 3, 12)


def h4(doc, text):
    return _heading(doc, text, 4, 12)


def caption(doc, text):
    p = para(doc, text, cn='宋体', size=10.5, bold=True, align=CENTER, indent=False,
             space_after=2, space_before=6)
    p.paragraph_format.keep_with_next = True  # 表题与表格不分离
    return p


def make_table(doc, headers, rows, widths=None, font_size=10):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, htext in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = CENTER
        r = p.add_run(htext)
        set_font(r, cn='黑体', size=font_size, bold=True)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.rows[i + 1].cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            r = p.add_run(str(val))
            set_font(r, cn='宋体', size=font_size)
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Cm(w)
    return t


def add_figure(doc, name, cap, width_cm=11.5):
    name = name.split('/')[-1]  # 归一化：content 中路径可能带 figures/ 前缀
    path = os.path.join(FIG_DIR, name)
    if not os.path.exists(path):
        para(doc, f'〔图示待渲染：{cap}〕', cn='宋体', size=10.5, align=CENTER,
             indent=False, color='888888', space_after=6, space_before=10)
        return
    p = doc.add_paragraph()
    p.alignment = CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True  # 图与图题不分离
    p.add_run().add_picture(path, width=Cm(width_cm))
    caption(doc, cap)


def add_toc(doc):
    p = doc.add_paragraph()
    r1 = p.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin'); r1._r.append(f1)
    r2 = p.add_run()
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve')
    it.text = 'TOC \\o "1-3" \\h \\z \\u'; r2._r.append(it)
    r3 = p.add_run()
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'separate'); r3._r.append(f2)
    r4 = p.add_run('（目录域：生成后需在 Word 中更新域显示）')
    set_font(r4, size=10.5)
    f3 = OxmlElement('w:fldChar'); f3.set(qn('w:fldCharType'), 'end'); r4._r.append(f3)


def add_page_number(doc):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = CENTER
    r = p.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin'); r._r.append(f1)
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = 'PAGE'; r._r.append(it)
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end'); r._r.append(f2)
    set_font(r, size=9)


def build():
    doc = Document()
    setup_styles(doc)
    # A4 页面 + 2.5cm 左右边距（压缩版式，兼顾打印规范）
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)
    doc.core_properties.title = '期末周别慌！系统需求分析'
    doc.core_properties.author = '刘添屹小组'

    # ---------- 封面 ----------
    para(doc, '', space_after=0)
    para(doc, f"文档编号：{META['doc_no']}", cn='宋体', size=10.5, indent=False, space_after=60)
    para(doc, META['title'], cn='黑体', size=30, bold=True, align=CENTER, indent=False, space_after=8)
    para(doc, META['subtitle'], cn='黑体', size=24, bold=True, align=CENTER, indent=False, space_after=40)
    para(doc, META['date'], cn='宋体', size=15, align=CENTER, indent=False, space_after=40)
    rows = [
        ('项目承担部门', META['dept']),
        ('项目组成员', META['members']),
        ('指导老师', META['advisor']),
        ('文档名称', META['doc_name']),
        ('文档类别', META['doc_type']),
        ('编制', META['writer']),
        ('编制时间', META['write_date']),
        ('校对', META['proofreader']),
        ('校对时间', META['proof_date']),
        ('评审负责人', META['reviewer']),
        ('评审时间', META['review_date']),
        ('批准', META['approver']),
        ('批准时间', META['approve_date']),
    ]
    t = make_table(doc, ['项目', '内容'], rows, widths=[4, 11], font_size=10.5)
    doc.add_paragraph()
    para(doc, META['division'], cn='宋体', size=10.5, indent=False, space_after=0)
    doc.add_page_break()

    # ---------- 修改记录（与目录同页起排，节省版式） ----------
    h1(doc, '修改记录')
    make_table(doc, ['版本号', '修改内容', '修改人', '备注', '修改日期'],
               [list(r) for r in REVISIONS], widths=[2, 5, 3, 2.5, 2.5])

    # ---------- 目录 ----------
    h1(doc, '目录')
    add_toc(doc)
    doc.add_page_break()

    # ---------- 1 引言 ----------
    h1(doc, '1 引言')
    h2(doc, '1.1 编写目的')
    para(doc, PURPOSE)
    para(doc, '本文档的预期读者包括：项目开发小组全体成员、指导老师与课程评审老师。')

    h2(doc, '1.2 项目背景')
    for b in BACKGROUND:
        para(doc, b)

    h2(doc, '1.3 定义')
    caption(doc, '表 1.3-1 名词定义表')
    make_table(doc, ['名词', '说明'], [list(d) for d in DEFINITIONS], widths=[3.5, 11])

    h2(doc, '1.4 参考资料')
    for i, ref in enumerate(REFERENCES, 1):
        para(doc, f'[{i}] {ref}', indent=False)

    # ---------- 2 任务概述 ----------
    h1(doc, '2 任务概述')
    h2(doc, '2.1 系统目标')
    for g in GOAL_PARAS:
        para(doc, g)

    h2(doc, '2.2 系统运行构架')
    h3(doc, '2.2.1 构架设计说明')
    for a in ARCH_PARAS:
        para(doc, a)
    add_figure(doc, 'fig_2_2_1.png', '图 2.2-1 系统总体构架图', width_cm=12)

    h3(doc, '2.2.2 运行环境')
    caption(doc, '表 2.2-1 运行环境定义表')
    make_table(doc, ['项目', '说明'], [list(r) for r in ENV_ROWS], widths=[4, 11])

    h2(doc, '2.3 外部角色与系统边界')
    para(doc, '系统外部角色定义见表 2.3-1。除学生、游客、管理员三类系统使用者外，DeepSeek 大模型 API 与'
              '四川大学教务处网站作为外部系统参与业务，运维人员由开发小组兼任。')
    caption(doc, '表 2.3-1 外部角色定义表')
    make_table(doc, ['角色', '说明'], [list(r) for r in ACTOR_ROWS], widths=[3.5, 11])
    para(doc, '系统边界说明如下。')
    for b in BOUNDARY_PARAS:
        para(doc, b)

    h2(doc, '2.4 隐藏需求分析')
    para(doc, '对照用户需求列表分析发现，以下需求在原始列表中未显式提出，但对系统正常运行必不可少，'
              '属于隐藏需求，本阶段予以补充（见表 2.4-1）。')
    caption(doc, '表 2.4-1 隐藏需求分析表')
    make_table(doc, ['隐藏需求', '问题分析', '处理方案'],
               [list(r) for r in HIDDEN_ROWS], widths=[3, 5.5, 6])

    # ---------- 3 功能需求 ----------
    h1(doc, '3 功能需求')

    h2(doc, '3.1 功能划分')
    para(doc, '系统功能需求划分为公共功能与八大业务模块：账号管理、考试与日程管理、复习资料管理、AI 智能分析、'
              '学习统计、提醒通知、特色功能、管理端功能。公共功能为多个业务模块复用的通用功能，单独提出。'
              '系统顶层用例图见图 3.1-1。')
    add_figure(doc, 'fig_3_1_1.png', '图 3.1-1 系统顶层用例图', width_cm=12)

    h2(doc, '3.2 需求分配')
    para(doc, '全部功能需求的分配见表 3.2-1，共 57 项，实现方式均为网页。优先级为"高"的功能是期末演示的主线，'
              '优先完成；"中"的功能在主线完成后完成；"低"的功能为扩展方向，按进度排期。')
    caption(doc, '表 3.2-1 系统需求分配表')
    make_table(doc, ['序号', '功能编号', '功能描述', '实现方式', '备注'],
               ALLOC_ROWS, widths=[1.2, 3.6, 5.2, 1.8, 3.2], font_size=8.5)

    h2(doc, '3.3 功能描述')
    para(doc, '以下按模块逐条描述各功能需求。每个功能包括：用例描述、前置条件、后置条件、参与者、输入数据、'
              '输出数据、事件流（活动图）等要素。各模块的事件流与活动图均以公共功能（身份认证、数据隔离等）'
              '为前提，不再重复描述。')

    for mi, mod in enumerate(MODULES):
        h3(doc, f"3.3.{mod['id']} {mod['title']}")
        para(doc, mod['intro'])
        add_figure(doc, mod['fig'], mod['figcap'], width_cm=12)
        for n, uc in enumerate(mod['usecases'], 1):
            h4(doc, f"3.3.{mod['id']}.{n} {uc['name']}（{uc['uid']}）")
            para(doc, f"用例描述：{uc['desc']}")
            t = make_table(doc, ['项目', '内容'], [
                ('前置条件', uc['pre']),
                ('后置条件', uc['post']),
                ('参与者', uc['actor']),
                ('输入数据', uc['inp']),
                ('输出数据', uc['out']),
            ], widths=[2.5, 12])
            add_figure(doc, uc['fig'], uc['figcap'], width_cm=10.5)

    # ---------- 4 数据描述 ----------
    h1(doc, '4 数据描述')
    h2(doc, '4.1 数据词典')
    para(doc, '系统核心数据的定义见以下数据词典表。')
    for title, name, desc, rows in DICT_TABLES:
        caption(doc, title)
        para(doc, f'标识符：{name}', indent=False, space_after=2, size=10.5)
        para(doc, f'描述：{desc}', indent=False, space_after=4, size=10.5)
        make_table(doc, ['数据项', '类型', '单位', '范围', '缺省值', '说明'],
                   [list(r) for r in rows], widths=[2.6, 1.8, 1.5, 3.6, 1.8, 3.7], font_size=8.5)

    h2(doc, '4.2 数据库描述')
    for p in DB_PARAS:
        para(doc, p, indent=False)

    # ---------- 5 性能需求 ----------
    h1(doc, '5 性能需求')
    h2(doc, '5.1 数据精确度')
    for p in NONFUNC['5.1']:
        para(doc, p)
    h2(doc, '5.2 时间特性')
    for p in NONFUNC['5.2']:
        para(doc, p)
    h2(doc, '5.3 适应性')
    for p in NONFUNC['5.3']:
        para(doc, p)

    # ---------- 6 产品质量需求 ----------
    h1(doc, '6 产品质量需求')
    h2(doc, '6.1 故障分析')
    for p in NONFUNC['6.1']:
        para(doc, p)
    h2(doc, '6.2 可靠性')
    for p in NONFUNC['6.2']:
        para(doc, p)

    # ---------- 7 其他需求 ----------
    h1(doc, '7 其他需求')
    h2(doc, '7.1 安全性')
    for p in NONFUNC['7.1']:
        para(doc, p)
    h2(doc, '7.2 易用性')
    for p in NONFUNC['7.2']:
        para(doc, p)
    h2(doc, '7.3 可维护性')
    for p in NONFUNC['7.3']:
        para(doc, p)
    h2(doc, '7.4 移植性')
    for p in NONFUNC['7.4']:
        para(doc, p)

    # ---------- 8 对照表 ----------
    h1(doc, '8 用户需求与系统需求规格对照表')
    para(doc, '将用户需求列表（FMZH-USR-0001~0067）逐条与本文档的'
              '系统需求规格对应，见表 8-1。隐藏需求补充项在"用户需求"列标注为"（隐藏需求）"。')
    caption(doc, '表 8-1 用户需求与系统需求规格对照表')
    make_table(doc, ['序号', '用户需求编号', '用户需求内容', '系统需求编号', '系统需求内容', '备注'],
               [list(r) for r in TRACE_ROWS],
               widths=[1.0, 2.6, 2.9, 2.8, 3.0, 2.7], font_size=8.5)

    # ---------- 9 其他 ----------
    h1(doc, '9 其他')
    for p in OTHER_PARAS:
        para(doc, p)

    add_page_number(doc)
    out = os.path.abspath(OUT)
    doc.save(out)
    print('已生成:', out)


if __name__ == '__main__':
    build()
