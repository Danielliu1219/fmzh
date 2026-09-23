# -*- coding: utf-8 -*-
"""生成《期末周别慌！》项目简要用户需求列表.docx，并转换为 .doc（交付格式）

需求编号已连续化：FMZH-USR-0001 ~ 0067（原列表 0019 缺号，全部后续编号提前一位）。
运行：python -X utf8 build_user_needs.py
"""
import os

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, '..', '..'))
DOCX_OUT = os.path.join(ROOT, '简要用户需求列表_刘添屹小组.docx')
DOC_OUT = os.path.join(ROOT, '简要用户需求列表_刘添屹小组.doc')

TITLE = '《期末周别慌！》项目简要用户需求列表'
WRITER = '撰写：刘添屹（组长）、许丽媛、胡歆桐、张诗琪、宋跃月'
INTRO = ('项目说明：本项目是面向大学生期末复习场景的一站式智能备考网站，集考试安排、复习计划、资料管理、'
         '大模型辅助分析、个性化复习推荐、学习统计和提醒通知于一体。系统重点通过整合历年题、教师重点、'
         '课堂资料、作业练习等多源信息，调用大模型分析常考必考内容、知识点权重，并根据不同学生目标和基础'
         '生成差异化复习方案。')

# FMZH-USR-0001 ~ 0067（连续编号）
ITEMS = [
    '学生可注册新账号，注册时需提交账号、密码、姓名、学号、专业等信息。',
    '系统应校验注册账号的唯一性，当账号已存在时提示用户重新输入。',
    '学生可通过账号和密码登录系统。',
    '学生连续三次输入错误密码后，系统应锁定该账号并提示用户进行找回或等待解锁。',
    '学生可编辑个人资料，包括头像、姓名、专业、年级、班级等信息。',
    '学生可注销个人账号，系统应在确认后处理相关用户数据。',
    '学生登录后，系统应展示“今日待办”聚合视图，包括考试、任务、提醒和复习进度。',
    '学生可录入个人考试安排，包括科目、时间、地点、时长等信息。',
    '学生可修改或删除已录入的个人考试安排。',
    '学生可一键导入由教师或管理员发布的统一考试安排到个人时间表。',
    '系统自动生成考试倒计时并在首页展示。',
    '学生可按科目或日期筛选查看自己的考试列表。',
    '学生可为各科目创建备考计划，并添加学习任务。',
    '学习任务应包含名称、科目、起止时间、地点、优先级、难度等信息。',
    '学生可编辑或删除已创建的学习任务。',
    '学习任务支持设置重复规则，包括每天、每周和自定义。',
    '系统可根据考试时间倒推自动生成备考时间表。',
    '系统应自动检测同一时段内的多个任务冲突并提示用户。',
    '学生可拖拽调整任务时间，并支持日视图、周视图、月视图切换。',
    '学生可为单个任务设置提醒时间。',
    '系统应提供番茄钟计时功能帮助学生专注学习。',
    '学生可设定每日学习时长目标，系统统计达成情况。',
    '学生可为各科目设置颜色标签进行区分。',
    '系统支持外部日历 ICS 导入和导出。',
    '学生可上传复习资料，支持 PDF、Word、PPT、TXT、图片等格式。',
    '系统应校验上传资料的文件格式、大小和安全性。',
    '学生上传资料时可添加标题、科目、标签和描述信息。',
    '学生可编辑、删除或移动分类已上传的复习资料。',
    '学生可上传多种类型的复习信息，包括历年试题、老师标注重点、平时作业、课堂笔记、教材章节、实验报告和练习题。',
    '系统应将同一科目的多份资料进行统一整理，形成该科目的综合复习资料库。',
    '系统可识别不同资料来源，并区分教师重点、历年题、课堂笔记、作业练习等资料类型。',
    '系统可对复习资料进行全文搜索，并按科目、标签、日期和资料类型筛选排序。',
    '学生可将重点、错题或高频题标记到“错题本/收藏”。',
    '学生可将多份资料合并为一份复习笔记，并导出为 PDF 或 Markdown。',
    '系统调用大模型自动提取复习资料中的核心知识点、摘要和复习提纲。',
    '系统调用大模型生成思维导图、知识框架、记忆卡片和背诵要点。',
    '学生可针对资料向大模型提问并获得解答，系统保存对话历史并支持多轮追问。',
    '系统调用大模型根据资料生成自测练习题。',
    '系统应基于多来源复习资料综合分析常考题型、必考知识点和高频考点。',
    '系统应根据历年题出现频率、教师强调次数、作业覆盖情况和资料标注信息，计算各知识点考试权重。',
    '系统应以可视化方式展示每个知识点的重要程度、掌握状态和建议投入时间。',
    '系统应允许学生选择备考目标，例如“保及格”“稳中等”“冲高分”等。',
    '系统应根据学生的备考目标、剩余时间、历史学习记录和自测结果，生成个性化知识点权重。',
    '对于基础薄弱或只求及格的学生，系统应优先推荐高频基础题、必考概念和短期提分内容。',
    '对于基础较好且追求高分的学生，系统应推荐综合题、难题、易失分点和拓展考点。',
    '系统应根据用户完成任务情况和自测表现动态调整复习计划和知识点权重。',
    '系统应解释个性化推荐依据，例如该知识点来自历年高频题、教师重点或用户薄弱项。',
    '系统应限制大模型输出范围，使其回答尽量基于用户上传资料、课程资料和考试目标。',
    '系统应提示大模型生成内容仅供复习参考，用户需结合教材和课堂要求判断。',
    '学生可对学习任务打卡标记完成。',
    '系统自动统计各科目学习时长，并以图表展示学习进度和任务完成率。',
    '系统生成备考周报或学习总结，结合大模型评测给出各科目掌握程度评分。',
    '学生可查看历史学习记录统计报表。',
    '系统在考试前和任务开始前发送站内或邮件提醒。',
    '学生可自定义提醒方式和提醒时间。',
    '学生可查看通知历史，并标记已读、未读或删除。',
    '系统可生成“期末压力指数”，综合考试密度、剩余任务量和完成率提示学生调整计划。',
    '系统可提供“临时救急模式”，在考试前短时间内生成高优先级冲刺复习清单。',
    '系统可根据错题本、收藏内容和高权重知识点生成考前重点回顾列表。',
    '系统可提供“今日最重要的三件事”功能，帮助学生降低期末周任务焦虑。',
    '系统可根据学生学习习惯推荐适合的学习时间段。',
    '系统可提供轻量级情绪记录功能，并根据压力状态给出复习节奏建议。',
    '管理员可发布统一考试安排，并可维护课程、科目和基础考试信息。',
    '管理员可管理异常账号和违规内容。',
    '所有用户数据应加密存储，普通学生只能访问和管理自己的个人数据。',
    '系统应记录关键操作日志，便于问题排查。',
    '系统应在删除账号、删除资料等重要操作前进行二次确认。',
]


def set_font(run, cn='宋体', size=12, bold=False):
    run.font.name = 'Times New Roman'
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn('w:eastAsia'), cn)
    run.font.size = Pt(size)
    run.font.bold = bold


def build_docx():
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
    normal.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(TITLE), cn='黑体', size=16, bold=True)

    p = doc.add_paragraph()
    set_font(p.add_run(WRITER), size=10.5)

    p = doc.add_paragraph()
    set_font(p.add_run(INTRO), size=12)

    for i, item in enumerate(ITEMS, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(f'FMZH-USR-{i:04d}'), cn='黑体', size=11, bold=True)
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(2)
        set_font(p2.add_run(item), size=12)

    doc.core_properties.title = TITLE
    doc.core_properties.author = '刘添屹小组'
    doc.save(DOCX_OUT)
    print('已生成:', DOCX_OUT)


def convert_to_doc():
    """用 Word COM 转换为 .doc（课程交付格式），先生成临时文件再原子替换。"""
    import win32com.client
    tmp = DOC_OUT + '.tmp'
    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    try:
        d = word.Documents.Open(DOCX_OUT)
        d.SaveAs2(tmp, FileFormat=0)  # 0 = wdFormatDocument97
        d.Close()
        if os.path.getsize(tmp) > 1000:
            os.replace(tmp, DOC_OUT)
            print('已生成:', DOC_OUT)
        else:
            print('转换结果异常，保留原 .doc 不动')
    finally:
        word.Quit()


if __name__ == '__main__':
    build_docx()
    convert_to_doc()
